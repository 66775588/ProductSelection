"""Excel / Word 导出：程序化构建多 sheet 数据表与结构化文档。"""
from __future__ import annotations

from pathlib import Path
from typing import Any


def export_excel(ctx: dict[str, Any], path: Path) -> None:
    """导出多 sheet Excel：结论 / 竞品 / 痛点 / 成本 / 风险。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()

    def _sheet(title, headers, rows):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for r in rows:
            ws.append(r)
        # 表头加粗
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 18
        return ws

    # 结论总览
    _sheet(
        "结论总览",
        ["项目", "内容"],
        [
            ["品类", ctx["category"]],
            ["平台", ctx["platform"]],
            ["推荐等级", ctx["recommendation"]],
            ["综合得分", ctx["total_score"]],
            ["判定依据", ctx["reason"]],
            ["数据来源", "；".join(f"{k}={v}" for k, v in ctx["provenance"].items())],
        ],
    )

    # 竞品
    _sheet(
        "竞品TOP10",
        ["排名", "商品", "价格", "月销", "店铺等级", "SKU数", "标签"],
        [
            [c["rank"], c["title"], c["price"], c["sales"], c["shop_level"], c["sku_count"], "/".join(c["tags"])]
            for c in ctx["competitors"]
        ],
    )

    # 痛点
    _sheet(
        "用户痛点",
        ["痛点", "频次", "严重度", "代表评论"],
        [[p["keyword"], p["frequency"], p["severity"], p["sample"]] for p in ctx["pain_points"]],
    )

    # 成本利润
    profit = ctx.get("profit") or {}
    if profit.get("cost_breakdown"):
        _sheet(
            "成本利润",
            ["成本项", "金额(元)"],
            [[k, v] for k, v in profit["cost_breakdown"].items()]
            + [
                ["毛利率(%)", profit.get("gross_margin")],
                ["月利润(元)", profit.get("monthly_profit")],
                ["回本周期(月)", profit.get("payback_months")],
                ["盈亏平衡月销", profit.get("break_even_units")],
            ],
        )

    # 风险
    _sheet(
        "风险清单",
        ["类型", "等级", "描述", "规避方案"],
        [[r["type"], r["level"], r["desc"], r["mitigation"]] for r in ctx["risk_items"]],
    )

    # 移除默认 sheet
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    wb.save(path)


def export_word(ctx: dict[str, Any], path: Path) -> None:
    """导出 Word 结构化报告。"""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(f"《{ctx['category']}》选品调研报告", level=0)
    doc.add_paragraph(f"平台：{ctx['platform']} ｜ 生成时间：{ctx['generated_at']}")

    # 一、结论
    doc.add_heading("一、选品结论总览", level=1)
    doc.add_paragraph(f"推荐等级：{ctx['recommendation']}    综合得分：{ctx['total_score']}/100")
    doc.add_paragraph(f"判定依据：{ctx['reason']}")
    doc.add_paragraph("核心优势：")
    for s in ctx["core_strengths"]:
        doc.add_paragraph(s, style="List Bullet")
    doc.add_paragraph("核心风险：")
    for r in ctx["core_risks"]:
        doc.add_paragraph(r, style="List Bullet")

    # 二、市场
    doc.add_heading("二、市场数据分析", level=1)
    market = ctx.get("market")
    if market:
        doc.add_paragraph(
            f"搜索量 {market['search_volume']}（{market['market_type']}，{market['seasonality']}）；"
            f"供需比 {market['supply_demand_ratio']}；转化率 {market['conversion_rate']}"
        )

    # 三、竞品
    doc.add_heading("三、竞品对标分析", level=1)
    ca = ctx.get("competitor_analysis") or {}
    if ca:
        doc.add_paragraph(
            f"竞品 {ca.get('count')} 个，价格带 {ca.get('price_band')}，"
            f"均价 {ca.get('avg_price')} 元，头部垄断率 {round(ca.get('top3_monopoly_rate',0)*100,1)}%"
        )
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "商品", "价格", "月销", "标签"
    for c in ctx["competitors"]:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text, row[3].text = (
            c["title"], str(c["price"]), str(c["sales"]), "/".join(c["tags"])
        )

    # 四、痛点
    doc.add_heading("四、用户痛点分析", level=1)
    for p in ctx["pain_points"]:
        doc.add_paragraph(f"{p['keyword']}（频次{p['frequency']}，{p['severity']}）", style="List Bullet")
    doc.add_paragraph("差异化方向：")
    for d in ctx["differentiation"]:
        doc.add_paragraph(d, style="List Bullet")

    # 五、利润
    doc.add_heading("五、成本利润测算", level=1)
    profit = ctx.get("profit") or {}
    if profit.get("cost_breakdown"):
        doc.add_paragraph(
            f"目标售价 {profit.get('selling_price')} 元，毛利率 {profit.get('gross_margin')}%，"
            f"回本周期 {profit.get('payback_months')} 个月"
        )

    # 六、风险
    doc.add_heading("六、风险排查清单", level=1)
    for r in ctx["risk_items"]:
        doc.add_paragraph(f"[{r['type']}/{r['level']}] {r['desc']} → {r['mitigation']}", style="List Bullet")

    # 七、落地
    doc.add_heading("七、落地建议", level=1)
    landing = ctx.get("landing") or {}
    doc.add_paragraph(f"切入价格带：{landing.get('price_band')}，建议目标售价 {landing.get('target_price')} 元")
    for s in landing.get("core_selling_points", []):
        doc.add_paragraph(s, style="List Bullet")

    doc.add_paragraph(
        "本报告由「电商选品调研 Agent」自动生成。标注「模拟数据」的指标为基于品类估算，非真实抓取。"
    ).runs[0].font.size = Pt(8)
    doc.save(path)
